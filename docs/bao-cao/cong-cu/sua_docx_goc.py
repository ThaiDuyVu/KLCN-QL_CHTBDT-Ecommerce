# -*- coding: utf-8 -*-
"""Áp dụng thay đổi (bản gốc txt -> bản sao đã chỉnh) trực tiếp lên file .docx tải từ Google Docs, giữ định dạng và hình."""
import copy
import difflib
import json
import os
import re
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from common import has_drawing, items_of, norm, rpr_set, set_update_fields

DOCX_IN, ORIG_TXT, EDIT_TXT, PNG_DIR, DOCX_OUT, REPORT = sys.argv[1:7]
doc = Document(DOCX_IN)
body = doc.element.body
W = lambda t: qn('w:' + t)
log = {'inplace': 0, 'inserted': 0, 'deleted_paras': 0, 'deleted_tables': 0, 'deleted_drawings': 0,
       'tables_new': 0, 'images_replaced': [], 'images_added': [], 'skipped_toc_ops': 0, 'notes': []}

# ------------------------------------------------------------------ đọc dữ liệu
def load(path):
    L = open(path, 'rb').read()[3:].decode('utf-8').split('\r\n')
    out = []
    for x in L:
        if 'Cogover Blog Hình 7.' in x:
            a, b = x.split('Hình 7.', 1)
            out += [a, 'Hình 7.' + b]
        else:
            out.append(x)
    return [(x, norm(x)) for x in out if norm(x) and norm(x) != '________________']


A = load(ORIG_TXT)
E = load(EDIT_TXT)
items = items_of(body)
segs = [x[2].split('\n') for x in items]
B = [((j, k), norm(s)) for j in range(len(items)) for k, s in enumerate(segs[j]) if norm(s)]
sm = difflib.SequenceMatcher(None, [a for _, a in A], [b for _, b in B], autojunk=False)
amap = {}
for tag, i1, i2, j1, j2 in sm.get_opcodes():
    if tag == 'equal':
        for k in range(i2 - i1):
            amap[i1 + k] = B[j1 + k][0]
el_index = {id(x[1]): j for j, x in enumerate(items)}

# ------------------------------------------------------------------ phân loại dòng
FRONT = {'LỜI CẢM ƠN', 'LỜI CAM ĐOAN', 'TÓM TẮT KHÓA LUẬN', 'MỤC LỤC', 'DANH MỤC HÌNH ẢNH', 'DANH MỤC BẢNG BIỂU',
         'DANH MỤC TỪ VIẾT TẮT', 'TÀI LIỆU THAM KHẢO', 'PHỤ LỤC'}


def body_start(L):
    return next(i for i, (_, n) in enumerate(L) if n == 'CHƯƠNG 1: GIỚI THIỆU & TỔNG QUAN ĐỀ TÀI')


A_BODY, E_BODY = body_start(A), body_start(E)


def kind(raw, front):
    s = norm(raw)
    if front:
        return 'h1' if s in FRONT else 'normal'
    if re.match(r'^CHƯƠNG \d+:', s) or s in FRONT:
        return 'h1'
    if s.startswith('[CẦN') or s.startswith('[GHI CHÚ'):
        return 'marker'
    if re.match(r'^\[\d+\]', s):
        return 'ref'
    if re.match(r'^Hình \d+(\.\d+|\.x)?\s*\.', s):
        return 'figcap'
    if re.match(r'^Bảng \d+\.(\d+|x)\.', s):
        return 'tabcap'
    if re.match(r'^UC(-SYS)?-\d\d\s*[-–]\s', s):
        return 'uctitle'
    if len(s) < 170 and re.match(r'^\d\.\d+\.\d+\.\d+\.\s*\S', s):
        return 'h4'
    if len(s) < 170 and re.match(r'^\d\.\d+\.\d+\.\s*\S', s):
        return 'h3'
    if len(s) < 170 and re.match(r'^\d\.\d+\.\s*\S', s):
        return 'h2'
    if '\t' in raw and not raw.startswith('\t'):
        return 'row'
    if re.match(r'^\s*[*•]\s', raw):
        return 'bullet'
    return 'normal'


# ------------------------------------------------------------------ mẫu định dạng lấy từ chính tài liệu
def first_text_run(p):
    for r in p.iter(W('r')):
        if r.find(W('t')) is not None:
            return r
    return None


def find_para(pattern, kind_filter=None, start=0, need_num=None):
    for j in range(start, len(items)):
        kd, p, t = items[j][0], items[j][1], items[j][2]
        if kind_filter and kd != kind_filter:
            continue
        if need_num is not None and (p.find(W('pPr') + '/' + W('numPr')) is not None) != need_num:
            continue
        if re.match(pattern, t):
            return p
    raise KeyError(pattern)


def a_para_of(pred):
    for i, (raw, n) in enumerate(A):
        if i >= A_BODY and pred(raw) and i in amap:
            j, _ = amap[i]
            if items[j][0] == 'p':
                return items[j][1]
    raise KeyError('bullet')


body_ch3 = el_index[id(find_para(r'^CHƯƠNG 3:'))]
TEMPLATES = {
    'h1': find_para(r'^CHƯƠNG 3:'),
    'h2': find_para(r'^3\.1\. ', 'p'),
    'h3': find_para(r'^3\.1\.1\. ', 'p'),
    'h4': find_para(r'^3\.1\.1\. ', 'p'),
    'normal': find_para(r'^Dựa trên kết quả khảo sát', 'p'),
    'bullet': next(items[amap[i][0]][1] for i in range(A_BODY, len(A)) if A[i][0].startswith('* ') and i in amap
                   and items[amap[i][0]][0] == 'p' and items[amap[i][0]][1].find(W('pPr') + '/' + W('numPr')) is not None),
    'figcap': find_para(r'^Hình 76', 'p'),
    'tabcap': find_para(r'^Bảng 2\.2', 'p'),
    'uctitle': find_para(r'^UC-01 - ', 'p'),
    'ref': find_para(r'^\[1\]', 'p', start=body_ch3),
}
TEMPLATES['marker'] = TEMPLATES['normal']
TPL = {k: (copy.deepcopy(p.find(W('pPr'))), copy.deepcopy(first_text_run(p).find(W('rPr'))) if first_text_run(p) is not None else None)
       for k, p in TEMPLATES.items()}
for k in TPL:  # bỏ sectPr (nếu có) khỏi mẫu
    ppr = TPL[k][0]
    if ppr is not None:
        for s in ppr.findall(W('sectPr')):
            ppr.remove(s)
cell_tpl_p = find_para(r'^Modular monolith', 'c')
CELL_TPL = (copy.deepcopy(cell_tpl_p.find(W('pPr'))), copy.deepcopy(first_text_run(cell_tpl_p).find(W('rPr'))))


def add_text(r, text):
    parts = text.split('\t')
    for k, part in enumerate(parts):
        if k:
            r.append(OxmlElement('w:tab'))
        if part:
            t = OxmlElement('w:t')
            t.set(qn('xml:space'), 'preserve')
            t.text = part
            r.append(t)


def make_run(rpr, text, highlight=False, bold=False, italic=False):
    r = OxmlElement('w:r')
    rp = copy.deepcopy(rpr) if rpr is not None else OxmlElement('w:rPr')
    if highlight:
        rpr_set(rp, 'highlight', 'yellow')
    if italic:
        rpr_set(rp, 'i')
    if bold:
        rpr_set(rp, 'b')
    r.append(rp)
    add_text(r, text)
    return r


def make_para(k, text):
    ppr, rpr = TPL[k]
    p = OxmlElement('w:p')
    if ppr is not None:
        p.append(copy.deepcopy(ppr))
    if k == 'normal' or k == 'marker' or k == 'bullet':
        s = norm(text) if k == 'bullet' else text.strip()
    else:
        s = text.strip()
    p.append(make_run(rpr, s, highlight=(k == 'marker'), italic=(k == 'marker')))
    return p


def set_segments(p, new_segs):
    """Đặt lại chữ của đoạn, giữ định dạng run đầu tiên, giữ hình/trường/ngắt trang."""
    first = first_text_run(p)
    rpr = copy.deepcopy(first.find(W('rPr'))) if first is not None and first.find(W('rPr')) is not None else None
    anchor_parent, anchor_pos = (first.getparent(), first.getparent().index(first)) if first is not None else (p, len(p))
    for r in list(p.iter(W('r'))):
        for ch in list(r):
            if ch.tag in (W('t'), W('tab')) or (ch.tag in (W('br'), W('cr')) and ch.get(W('type')) in (None, 'textWrapping')):
                r.remove(ch)
        if all(ch.tag == W('rPr') for ch in r):
            par = r.getparent()
            if par is anchor_parent and par.index(r) < anchor_pos:
                anchor_pos -= 1
            par.remove(r)
    nr = OxmlElement('w:r')
    if rpr is not None:
        nr.append(rpr)
    for k, sgt in enumerate(new_segs):
        if k:
            nr.append(OxmlElement('w:br'))
        add_text(nr, sgt.strip() if not sgt.startswith('\t') else sgt.lstrip('\t'))
    anchor_parent.insert(min(anchor_pos, len(anchor_parent)), nr)


def top_level(el):
    while el.getparent() is not body and el.getparent().tag != W('sdtContent'):
        el = el.getparent()
    return el


def tbl_of(j):
    return items[j][3]


# ------------------------------------------------------------------ bảng mới (từ dòng "cột<TAB>cột")
def make_table(rows):
    ncol = max(len(r) for r in rows)
    tbl = OxmlElement('w:tbl')
    tpr = OxmlElement('w:tblPr')
    tw = OxmlElement('w:tblW'); tw.set(W('w'), '5000'); tw.set(W('type'), 'pct'); tpr.append(tw)
    jc = OxmlElement('w:jc'); jc.set(W('val'), 'center'); tpr.append(jc)
    bd = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side); e.set(W('val'), 'single'); e.set(W('sz'), '4'); e.set(W('space'), '0'); e.set(W('color'), '000000')
        bd.append(e)
    tpr.append(bd)
    tbl.append(tpr)
    grid = OxmlElement('w:tblGrid')
    for _ in range(ncol):
        g = OxmlElement('w:gridCol'); g.set(W('w'), str(int(9070 / ncol))); grid.append(g)
    tbl.append(grid)
    for ri, r in enumerate(rows):
        tr = OxmlElement('w:tr')
        if ri == 0:
            trp = OxmlElement('w:trPr'); h = OxmlElement('w:tblHeader'); trp.append(h); tr.append(trp)
        for ci in range(ncol):
            tc = OxmlElement('w:tc')
            tcp = OxmlElement('w:tcPr')
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(W('val'), 'clear'); sh.set(W('color'), 'auto'); sh.set(W('fill'), 'D9E2F3'); tcp.append(sh)
            tc.append(tcp)
            p = OxmlElement('w:p')
            p.append(copy.deepcopy(CELL_TPL[0]))
            text = r[ci] if ci < len(r) else ''
            m = re.search(r'(\[(?:CẦN|GHI CHÚ)[^\]]*\])', text)
            p.append(make_run(CELL_TPL[1], text, bold=(ri == 0), highlight=bool(m)))
            tc.append(p)
            tr.append(tc)
        tbl.append(tr)
    log['tables_new'] += 1
    return tbl


# ------------------------------------------------------------------ vùng đặc biệt
def idx(L, text, start=0):
    return next(i for i in range(start, len(L)) if L[i][1] == text)


A_TOC = (idx(A, 'MỤC LỤC'), idx(A, 'DANH MỤC HÌNH ẢNH', idx(A, 'MỤC LỤC') + 1))
E_TOC = (idx(E, 'MỤC LỤC'), idx(E, 'DANH MỤC HÌNH ẢNH', idx(E, 'MỤC LỤC') + 1))
A_ABB = (idx(A, 'Từ viết tắt'), idx(A, 'Create - Read - Update - Delete'))
E_ABB = (idx(E, 'Từ viết tắt'), idx(E, 'XSS') + 1)

# ------------------------------------------------------------------ áp dụng thay đổi
newsegs = {}          # j -> danh sách đoạn (None = xóa)
deleted_items = set()
e_elem = {}           # chỉ số dòng E -> phần tử docx của nó
after_body = {}       # j -> phần tử body cuối cùng đã chèn sau phần tử j (giữ thứ tự)
seg_inserts = {}      # (j,k) -> danh sách dòng chèn thêm vào sau đoạn k của item j
pending_png = []


def cur_segs(j):
    if j not in newsegs:
        newsegs[j] = list(segs[j])
    return newsegs[j]


def ptxt(e):
    return norm(''.join(t.text or '' for t in e.iter(W('t'))))


def insert_body_after(j_or_el, el, after_images=False):
    base = j_or_el if not isinstance(j_or_el, int) else items[j_or_el][1]
    key = id(base)
    last = after_body.get(key, base)
    if after_images:  # chú thích hình phải nằm dưới ảnh
        nx = last.getnext()
        while nx is not None and nx.tag == W('p') and not ptxt(nx) and nx.find(W('pPr') + '/' + W('sectPr')) is None:
            last = nx
            nx = nx.getnext()
    last.addnext(el)
    after_body[key] = el
    return el


def sim(a, b):
    return difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()


def pair(aa, ee):
    """Ghép cặp có thứ tự giữa dòng cũ và dòng mới theo độ giống nhau (quy hoạch động)."""
    n, m = len(aa), len(ee)
    S = [[0.0] * (m + 1) for _ in range(n + 1)]
    for x in range(n - 1, -1, -1):
        for y in range(m - 1, -1, -1):
            r = sim(aa[x], ee[y])
            S[x][y] = max(S[x + 1][y], S[x][y + 1], (S[x + 1][y + 1] + r) if r > 0.45 else -1)
    res, x, y = [], 0, 0
    while x < n and y < m:
        r = sim(aa[x], ee[y])
        if r > 0.45 and abs(S[x][y] - (S[x + 1][y + 1] + r)) < 1e-9:
            res.append((x, y)); x += 1; y += 1
        elif S[x][y] == S[x + 1][y]:
            x += 1
        else:
            y += 1
    return res


def emit_new(e_list, cursor_j, cursor_seg, next_j):
    """Chèn các dòng mới e_list sau vị trí cursor."""
    rows = []

    def flush_rows():
        if rows:
            base = top_level(items[cursor_j][1]) if cursor_j is not None else None
            t = make_table([r.split('\t') for r in rows])
            insert_body_after(base, t)
            rows.clear()

    for ei in e_list:
        raw = E[ei][0]
        k = kind(raw, ei < E_BODY)
        if k == 'marker' and '.puml' in raw and ('VẼ LẠI' in raw or 'CHÈN HÌNH' in raw):
            pending_png.append(re.search(r'so-do/([\w\-]+)\.puml', raw).group(1))
            continue
        if k == 'row':
            rows.append(raw); continue
        flush_rows()
        in_cell = cursor_j is not None and items[cursor_j][0] == 'c'
        if in_cell and not (next_j is None or tbl_of(next_j) is not tbl_of(cursor_j)):
            seg_inserts.setdefault((cursor_j, cursor_seg), []).append(raw)  # còn trong bảng: thêm vào ô
            continue
        base = top_level(items[cursor_j][1]) if in_cell else items[cursor_j][1]
        el = insert_body_after(base, make_para(k if k != 'h4' else 'h3', raw), after_images=(k == 'figcap'))
        e_elem[ei] = el
        log['inserted'] += 1
        if k == 'figcap' and pending_png:
            log.setdefault('_png_for', []).append((pending_png.pop(0), el))
    flush_rows()


def mapped_before(i):
    while i >= 0 and i not in amap:
        i -= 1
    return amap.get(i)


def mapped_from(i):
    while i < len(A) and i not in amap:
        i += 1
    return amap.get(i)


def clean(raw):
    return re.sub(r'^\s*([*•]\s+)?', '', raw).rstrip()


def handle_pair(ax, ey):
    j, k = amap[ax]
    ak = kind(A[ax][0], ax < A_BODY)
    ek = kind(E[ey][0], ey < E_BODY)
    puml = ek == 'marker' and '.puml' in E[ey][0]
    if (items[j][0] == 'c' or ak == ek or {ak, ek} <= {'h3', 'h4'}) and ek != 'row' and not puml:
        cur_segs(j)[k] = clean(E[ey][0])
        e_elem[ey] = items[j][1]
        log['inplace'] += 1
        if ek == 'figcap' and pending_png:
            log.setdefault('_png_for', []).append((pending_png.pop(0), items[j][1]))
    else:
        cur_segs(j)[k] = None
        emit_new([ey], j, k, None)
    return (j, k)


ops = difflib.SequenceMatcher(None, [a for _, a in A], [e for _, e in E], autojunk=False).get_opcodes()
for tag, i1, i2, j1, j2 in ops:
    if tag == 'equal':
        for k in range(i2 - i1):
            if i1 + k in amap:
                e_elem[j1 + k] = items[amap[i1 + k][0]][1]
        continue
    if A_TOC[0] < i1 and i2 <= A_TOC[1]:
        log['skipped_toc_ops'] += 1
        continue
    if A_ABB[0] <= i1 and i2 <= A_ABB[1] + 1:
        continue  # bảng từ viết tắt dựng lại riêng
    aa, ee = list(range(i1, i2)), list(range(j1, j2))
    pairs = pair([A[x][1] for x in aa], [E[y][1] for y in ee]) if aa and ee else []
    pairs = [(aa[x], ee[y]) for x, y in pairs if aa[x] in amap]
    cursor = mapped_before(i1 - 1)
    op_next = mapped_from(i2)
    pa, pe = i1, j1
    for ax, ey in pairs + [(i2, j2)]:
        # xóa dòng cũ không ghép được
        for x in range(pa, ax):
            if x in amap:
                j, k = amap[x]
                cur_segs(j)[k] = None
        # chèn dòng mới không ghép được, sau con trỏ
        if pe < ey:
            nxt = amap[ax][0] if ax < i2 and ax in amap else (op_next[0] if op_next else None)
            if cursor is None:
                cursor = mapped_from(i1)
            emit_new(list(range(pe, ey)), cursor[0], cursor[1], nxt)
        if ax < i2:
            cursor = handle_pair(ax, ey)
        pa, pe = ax + 1, ey + 1

# Gắn dòng chèn trong ô bảng
for (j, k), extra in seg_inserts.items():
    s = cur_segs(j)
    pos = k + 1
    while pos < len(s) and s[pos] is not None and False:
        pos += 1
    for x in extra:
        s.insert(pos, x); pos += 1
        log['inserted'] += 1

# Viết lại các đoạn đã đổi
for j, s in newsegs.items():
    p = items[j][1]
    live = [x for x in s if x is not None]
    if not live:
        if p.find(W('pPr') + '/' + W('sectPr')) is not None or has_drawing(p):
            set_segments(p, [''])
            continue
        deleted_items.add(j)
        continue
    set_segments(p, live)

# Xóa đoạn; ô bảng phải còn ít nhất 1 đoạn; bảng bị xóa hết thì xóa cả bảng
tables_touched = {}
for j in sorted(deleted_items):
    kd, p, _, tbl, tc = items[j]
    if kd == 'c':
        tables_touched.setdefault(id(tbl), tbl)
    p.getparent().remove(p)
    log['deleted_paras'] += 1
for j, x in enumerate(items):
    if x[0] == 'c' and x[4] is not None and x[4].find(W('p')) is None:
        e = OxmlElement('w:p'); x[4].append(e)
for t in tables_touched.values():
    alive = [j for j, x in enumerate(items) if x[3] is t and j not in deleted_items and norm(x[2])]
    if not alive and t.getparent() is not None:
        t.getparent().remove(t)
        log['deleted_tables'] += 1

# Hình mồ côi: đoạn chỉ chứa hình nằm giữa hai đoạn chữ đã bị xóa
alive_text = [j for j, x in enumerate(items) if norm(x[2]) and j not in deleted_items]
deleted_text = [j for j in deleted_items if norm(items[j][2])]
for j, x in enumerate(items):
    if x[0] == 'p' and has_drawing(x[1]) and not norm(x[2]) and x[1].getparent() is not None:
        prev_t = max((a for a in alive_text + deleted_text if a < j), default=None)
        next_t = min((a for a in alive_text + deleted_text if a > j), default=None)
        if prev_t in deleted_items and next_t in deleted_items:
            x[1].getparent().remove(x[1])
            log['deleted_drawings'] += 1

# Bảng từ viết tắt: dựng lại theo bản sao
abb_tbl = next(x[3] for x in items if x[0] == 'c' and x[2].strip() == 'Từ viết tắt')
pairs_e = [E[k][1] for k in range(E_ABB[0], E_ABB[1] + 1)]
rows = abb_tbl.findall(W('tr'))
while len(rows) < len(pairs_e) // 2:
    nr = copy.deepcopy(rows[-1]); rows[-1].addnext(nr); rows.append(nr)
for r_i, tr in enumerate(rows):
    tcs = tr.findall(W('tc'))
    for c_i, tc in enumerate(tcs[:2]):
        p = tc.find(W('p'))
        set_segments(p, [pairs_e[r_i * 2 + c_i]])

# Ảnh sơ đồ PlantUML
sec = doc.sections[0]
width = sec.page_width - sec.left_margin - sec.right_margin
png_for = log.pop('_png_for', [])
# Hình 3.6 và 3.24: vẽ lại theo Giai đoạn 3 (ghi trong 03_CHUONG_3_YEU_CAU_USECASE.md, không có dòng đánh dấu trong bản sao)
for cap_text, png in (('Hình 3.6. Sơ đồ Use-Case nghiệp vụ tổng quát', 'hinh13_usecase_nghiep_vu_tong_quat'),
                      ('Hình 3.24. Sơ đồ Use-Case hệ thống tổng quát', 'hinh31_usecase_he_thong_tong_quat')):
    caps = [x[1] for x in items_of(body) if x[0] == 'p' and x[2].strip() == cap_text]
    png_for.append((png, caps[-1]))
seen = set()
for png, cap in png_for:
    if png in seen:
        log['notes'].append('giữ ảnh cũ (sơ đồ dùng chung): ' + png)
        continue
    seen.add(png)
    path = os.path.join(PNG_DIR, png + '.png')
    target = None
    if has_drawing(cap):
        target = cap
    else:
        # Tìm ảnh cũ ngay phía trên chú thích: bỏ qua đoạn trống; ảnh có thể nằm chung đoạn với chữ
        sib = cap.getprevious()
        steps = 0
        while sib is not None and sib.tag == W('p') and steps < 4:
            if has_drawing(sib):
                target = sib
                break
            st = sib.find(W('pPr') + '/' + W('pStyle'))
            if ptxt(sib).startswith('Hình') or (st is not None and st.get(W('val')).startswith('Heading')):
                break
            sib = sib.getprevious()
            steps += 1
    if target is not None and target is not cap and norm(''.join(t.text or '' for t in target.iter(W('t')))):
        # Ảnh dính chung đoạn văn: gỡ ảnh cũ, đặt ảnh mới vào đoạn riêng ngay trên chú thích
        for r in list(target.iter(W('r'))):
            if r.find(W('drawing')) is not None or r.find(W('pict')) is not None:
                r.getparent().remove(r)
        log['notes'].append('tách ảnh khỏi đoạn văn: ' + png)
        target = None
    if target is None:
        target = OxmlElement('w:p')
        ppr = OxmlElement('w:pPr'); jc = OxmlElement('w:jc'); jc.set(W('val'), 'center'); ppr.append(jc); target.append(ppr)
        cap.addprevious(target)
        log['images_added'].append(png)
    else:
        for r in list(target.iter(W('r'))):
            if r.find(W('drawing')) is not None or r.find(W('pict')) is not None:
                r.getparent().remove(r)
        log['images_replaced'].append(png)
    Paragraph(target, doc._body).add_run().add_picture(path, width=width)

# Hình 3.83: sơ đồ cũ, phần mới đã gộp vào Hình 3.82
for x in items_of(body):
    if x[0] == 'p' and x[2].strip().startswith('Hình 3.83.') and x[1].getprevious() is not None and has_drawing(x[1].getprevious()):
        x[1].addnext(make_para('marker', '[GHI CHÚ – ẢNH CŨ: sơ đồ lớp mới của phân hệ Chatbot và RAG đã gộp chung vào Hình 3.82; nên xóa ảnh này và gộp mục 3.7.4.7 vào 3.7.4.6, hoặc giữ ảnh cũ nếu muốn tách riêng]'))
        log['notes'].append('ghi chú Hình 3.83')

# Cập nhật trường (mục lục) khi mở file
set_update_fields(doc.settings.element)
doc.save(DOCX_OUT)
json.dump(log, open(REPORT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print({k: (len(v) if isinstance(v, list) else v) for k, v in log.items()})
