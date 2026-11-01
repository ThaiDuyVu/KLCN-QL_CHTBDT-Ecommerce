# -*- coding: utf-8 -*-
"""Chuyển BaoCaoKhoaLuan_ChinhSua.txt (bản xuất từ Word) thành .docx có style, bảng, mục lục tự động và sơ đồ đã render."""
import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC, FIGMAP, PNG_DIR, OUT = sys.argv[1:5]
L = open(SRC, 'rb').read()[3:].decode('utf-8').split('\r\n')
old_of = {f['new']: f['old'] for f in json.load(open(FIGMAP, encoding='utf-8'))['figures']}

FONT = 'Times New Roman'
doc = Document()

# ---------------------------------------------------------------- trang, font, style
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.orientation = WD_ORIENT.PORTRAIT
sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(2), Cm(2), Cm(3), Cm(2)
TEXT_WIDTH = Cm(16)


def set_font(style, size, bold=None, italic=None, color=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(a), FONT)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = color or RGBColor(0, 0, 0)


st = doc.styles
normal = st['Normal']
set_font(normal, 13)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.space_before = Pt(0)

for name, size, bold, italic, before, after in [
    ('Heading 1', 14, True, False, 0, 12),
    ('Heading 2', 13, True, False, 12, 6),
    ('Heading 3', 13, True, True, 10, 6),
    ('Heading 4', 13, False, True, 8, 4),
]:
    s = st[name]
    set_font(s, size, bold, italic)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.line_spacing = 1.3
st['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_style(name, size, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=False):
    s = st.add_style(name, 1)
    s.base_style = normal
    set_font(s, size, bold, italic)
    s.paragraph_format.alignment = align
    s.paragraph_format.line_spacing = 1.15
    s.paragraph_format.space_before = Pt(3)
    s.paragraph_format.space_after = Pt(9)
    s.paragraph_format.keep_with_next = keep_next
    return s


add_style('Caption Hinh', 12, italic=True)
add_style('Caption Bang', 12, bold=True, keep_next=True)
add_style('Front Heading', 14, bold=True)
add_style('UC Title', 13, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
add_style('Ghi chu', 11, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT)
add_style('Anh giu cho', 11, italic=True)
code = add_style('Code', 10, align=WD_ALIGN_PARAGRAPH.LEFT)
set_font(code, 10)
code.font.name = 'Consolas'
code.paragraph_format.space_after = Pt(0)
cell_style = add_style('Table Text', 11, align=WD_ALIGN_PARAGRAPH.LEFT)
cell_style.paragraph_format.space_before = Pt(1)
cell_style.paragraph_format.space_after = Pt(1)
cell_style.paragraph_format.line_spacing = 1.0

# Chân trang: số trang
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_field(par, instr, placeholder=''):
    r = par.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
    r = par.add_run()
    t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr; r._r.append(t)
    r = par.add_run()
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate'); r._r.append(s)
    par.add_run(placeholder)
    r = par.add_run()
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)


add_field(fp, 'PAGE', '1')
# Yêu cầu Word cập nhật trường (mục lục, danh mục) khi mở
upd = OxmlElement('w:updateFields'); upd.set(qn('w:val'), 'true')
doc.settings.element.append(upd)

# ---------------------------------------------------------------- tiện ích ghi
state = {'last_break': True}


def para(text='', style=None, align=None, bold=False):
    p = doc.add_paragraph(style=style)
    if text:
        run_text(p, text, bold)
    if align is not None:
        p.alignment = align
    state['last_break'] = False
    return p


MARK = re.compile(r'(\[(?:CẦN|GHI CHÚ)[^\]]*\])')


def run_text(p, text, bold=False):
    for part in MARK.split(text):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold or None
        if MARK.fullmatch(part):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.italic = True


def page_break():
    if state['last_break']:
        return
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    state['last_break'] = True


def placeholder(text):
    p = para(style='Anh giu cho')
    r = p.add_run(text)
    r.font.highlight_color = WD_COLOR_INDEX.TURQUOISE
    pbdr = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + side)
        el.set(qn('w:val'), 'dashed'); el.set(qn('w:sz'), '6'); el.set(qn('w:space'), '4'); el.set(qn('w:color'), '808080')
        pbdr.append(el)
    p._p.get_or_add_pPr().append(pbdr)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    tcpr.append(sh)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trpr.append(h)


def table(rows, header=True, label_col=False):
    """rows: list[list[list[str]]] — mỗi ô là danh sách đoạn văn."""
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 10 if ncol >= 5 else 11
    for ri, r in enumerate(rows):
        row = t.add_row()
        for ci in range(ncol):
            paras = r[ci] if ci < len(r) else ['']
            cell = row.cells[ci]
            for k, text in enumerate(paras):
                p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                p.style = st['Table Text']
                run_text(p, text, bold=(header and ri == 0) or (label_col and ci == 0))
                for run in p.runs:
                    run.font.size = Pt(size)
            if header and ri == 0:
                shade(cell, 'D9E2F3')
        if header and ri == 0:
            repeat_header(row)
    if label_col:
        for row in t.rows:
            row.cells[0].width = Cm(4.2)
            row.cells[1].width = TEXT_WIDTH - Cm(4.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    state['last_break'] = False


# ---------------------------------------------------------------- nhận diện dòng
H1 = re.compile(r'^CHƯƠNG \d+:')
H4 = re.compile(r'^\d\.\d+\.\d+\.\d+\.\s*\S')
H3 = re.compile(r'^\d\.\d+\.\d+\.\s*\S')
H2 = re.compile(r'^\d\.\d+\.\s*\S')
FIG = re.compile(r'^\s*(Hình (\d+)\.(\d+))\.\s')
TAB_CAP = re.compile(r'^\t?Bảng \d+\.\d+\.\s')
UC_TITLE = re.compile(r'^UC(-SYS)?-\d\d\s*[-–]\s')
LABEL = re.compile(r'^(Tên|Tên UseCase|Tên Use Case|Tên:|ID|Use Case|Mô tả ngắn gọn|Mô tả|Tác nhân( \(Actors\))?|Điều kiện tiên quyết.*|Điều kiện sau.*|Luồng sự kiện chính.*|Luồng phụ / Ngoại lệ|Luồng sự kiện phụ.*)$')
PUML = re.compile(r'docs/bao-cao/so-do/([\w\-]+)\.puml')
EQ = re.compile(r'^\s{8,}\((\d\.\d)\)\s*$')
FRONT_H1 = {'LỜI CẢM ƠN', 'LỜI CAM ĐOAN', 'TÓM TẮT KHÓA LUẬN', 'MỤC LỤC', 'DANH MỤC HÌNH ẢNH', 'DANH MỤC BẢNG BIỂU',
            'DANH MỤC TỪ VIẾT TẮT', 'TÀI LIỆU THAM KHẢO', 'PHỤ LỤC'}
FRONT_BOLD = {'NHIỆM VỤ KHÓA LUẬN TỐT NGHIỆP', 'NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN', 'NHẬN XÉT CỦA GIẢNG VIÊN PHẢN BIỆN',
              'PHẦN DÀNH CHO KHOA, BỘ MÔN:', 'KHÓA LUẬN TỐT NGHIỆP', 'XÂY DỰNG HỆ THỐNG QUẢN LÝ CỬA HÀNG', 'MUA BÁN THIẾT BỊ ĐIỆN TỬ'}
EXPORT_COLS = {'\tÝ nghĩa': 2, '\tModular monolith': 4, '\tPostgreSQL + pgvector': 3, '\tOllama (mô hình cục bộ)': 3,
               '\tCông cụ chính': 3, '\tPackage': 3}

body_start = next(i for i, l in enumerate(L) if l == 'CHƯƠNG 1: GIỚI THIỆU & TỔNG QUAN ĐỀ TÀI')
uc_s = next(i for i, l in enumerate(L) if l.startswith('3.3. Sơ đồ Use-Case') and i > body_start)
uc_e = next(i for i, l in enumerate(L) if l.strip().startswith('3.4. Sơ đồ hoạt động') and i > uc_s)

pending = {'png': None, 'used': None, 'marker': '', 'alt': False}
stats = {'png': 0, 'placeholder': 0, 'tables': 0}


def figure(caption_line):
    m = FIG.match(caption_line)
    num = m.group(1)
    if pending['png']:
        doc.add_picture(os.path.join(PNG_DIR, pending['png'] + '.png'), width=TEXT_WIDTH)
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
        pending['used'], pending['png'] = pending['png'], None
        stats['png'] += 1
    elif pending['used'] and num in pending['marker']:
        placeholder('[Dùng chung sơ đồ ở hình phía trên (%s.puml) – có thể gộp hai mục thành một hình]' % pending['used'])
    else:
        old = old_of.get(num, num.split()[-1] if num.startswith('Hình 2.') else None)
        if old is None or old.endswith('.x'):
            placeholder('[ẢNH %s – cần chụp màn hình / bổ sung]' % num)
        else:
            placeholder('[ẢNH %s – chép từ bản Word gốc (Hình %s)]' % (num, old.replace('Hình ', '')))
        stats['placeholder'] += 1
    pending['alt'] = False
    para(caption_line.strip(), style='Caption Hinh')


def flush_alt():
    if pending['alt']:
        placeholder('[ẢNH trong bản Word gốc (không có chú thích) – kiểm tra giữ hay bỏ]')
        pending['alt'] = False


i = 0
n = len(L)
while i < n:
    raw = L[i]
    line = raw.strip()
    front = i < body_start

    # ---- ngắt trang
    if line == '________________':
        flush_alt(); page_break(); i += 1; continue
    if line == '' or (raw and not line):
        i += 1; continue

    # ---- mục lục / danh mục: thay bằng trường tự động
    if front and line in ('MỤC LỤC', 'DANH MỤC HÌNH ẢNH', 'DANH MỤC BẢNG BIỂU'):
        page_break()
        para(line, style='Heading 1')
        p = para()
        instr = {'MỤC LỤC': 'TOC \\o "1-4" \\h \\z \\u',
                 'DANH MỤC HÌNH ẢNH': 'TOC \\h \\z \\t "Caption Hinh,1"',
                 'DANH MỤC BẢNG BIỂU': 'TOC \\h \\z \\t "Caption Bang,1"'}[line]
        add_field(p, instr, '[Nhấn chuột phải → Update Field để tạo %s]' % line.lower())
        i += 1
        while i < n and L[i].strip() != '________________':
            i += 1
        continue

    # ---- phần đầu
    if front:
        if line in FRONT_H1:
            page_break(); para(line, style='Heading 1')
        elif line in FRONT_BOLD:
            para(line, style='Front Heading')
        elif line.startswith('HUIT - Trường Đại học'):
            placeholder('[LOGO TRƯỜNG – chép từ bản Word gốc]')
        elif line == 'TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH' and i > 5:
            page_break(); para(line, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        elif i < 20 or line.startswith(('TRƯỜNG ĐẠI HỌC', 'KHOA CÔNG NGHỆ', '----', 'THÀNH PHỐ HỒ CHÍ MINH')):
            para(line, align=WD_ALIGN_PARAGRAPH.CENTER, bold=line.startswith(('TRƯỜNG', 'KHOA')))
        elif line == 'Từ viết tắt' and L[i + 1] == '\tÝ nghĩa':
            cells = [line] + [L[k].strip() for k in range(i + 1, n) if L[k].startswith('\t') and '____' not in L[k] and k < i + 200]
            j = i + 1
            while L[j].startswith('\t') and '____' not in L[j]:
                j += 1
            cells = [line] + [L[k].strip() for k in range(i + 1, j)]
            table([[[cells[k]], [cells[k + 1]]] for k in range(0, len(cells), 2)])
            stats['tables'] += 1
            i = j; continue
        else:
            para(line, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        i += 1; continue

    # ---- bảng xuất từ Word (ngoài mục Use Case): số cột biết trước
    if not raw.startswith('\t') and '\t' not in raw and i + 1 < n and L[i + 1] in EXPORT_COLS:
        nc = EXPORT_COLS[L[i + 1]]
        cells = [line]
        j = i + 1
        while j < n and L[j].startswith('\t') and '____' not in L[j]:
            if len(cells) % nc == 0 and L[j].strip().startswith(('Nguồn', 'Việc')):
                break
            cells.append(L[j].strip()); j += 1
        assert len(cells) % nc == 0, (i + 1, len(cells), nc)
        table([[[c] for c in cells[k:k + nc]] for k in range(0, len(cells), nc)])
        stats['tables'] += 1
        i = j; continue

    # ---- bảng đặc tả Use Case (2 cột: nhãn | nội dung)
    if uc_s < i < uc_e and not raw.startswith('\t') and LABEL.match(line):
        rows = []
        j = i
        while j < n:
            r, s = L[j], L[j].strip()
            # Ảnh trong bản gốc để lại dòng chỉ có dấu cách: đó là hết bảng. Dòng chỉ có TAB là ô trống.
            if j > i and (UC_TITLE.match(s) or H2.match(s) or H3.match(s) or FIG.match(r)
                          or (len(r) >= 2 and not s and '\t' not in r)):
                break
            if s == '':
                j += 1; continue
            if LABEL.match(s) and (not rows or rows[-1][1]):
                rows.append([[s], []])
            elif rows:
                rows[-1][1].append(s)
            j += 1
        hdr = rows and rows[0][0] == ['Use Case']
        table([[a, b or ['']] for a, b in rows], header=bool(hdr), label_col=True)
        stats['tables'] += 1
        i = j; continue

    # ---- bảng dạng "cột<TAB>cột" (bảng thêm ở Giai đoạn 5)
    if '\t' in raw and not raw.startswith('\t') and i + 1 < n and '\t' in L[i + 1] and not L[i + 1].startswith('\t'):
        j = i
        while j < n and '\t' in L[j] and not L[j].startswith('\t'):
            j += 1
        table([[[c] for c in L[k].split('\t')] for k in range(i, j)])
        stats['tables'] += 1
        i = j; continue

    # ---- công thức, mã nguồn
    m = EQ.match(raw)
    if m:
        placeholder('[CÔNG THỨC (%s) – chép từ bản Word gốc]' % m.group(1)); i += 1; continue
    if line == '{' and i + 1 < n and L[i + 1].startswith('  "'):
        j = i
        while j < n and L[j].strip() != '}':
            para(L[j], style='Code'); j += 1
        para('}', style='Code'); i = j + 1; continue

    # ---- ảnh mất chú thích (dòng alt text bắt đầu bằng dấu cách)
    if re.match(r'^ {1,2}\S', raw) and not (uc_s < i < uc_e):
        flush_alt(); pending['alt'] = True; i += 1; continue

    # ---- tiêu đề, chú thích, ghi chú
    s = line
    if H1.match(s):
        flush_alt(); page_break(); para(s, style='Heading 1'); i += 1; continue
    if s in FRONT_H1:
        flush_alt(); page_break(); para(s, style='Heading 1'); i += 1; continue
    if len(s) < 170 and H4.match(s):
        flush_alt(); para(s, style='Heading 4'); i += 1; continue
    if len(s) < 170 and H3.match(s):
        flush_alt(); para(s, style='Heading 3'); i += 1; continue
    if len(s) < 170 and H2.match(s):
        flush_alt(); para(s, style='Heading 2'); i += 1; continue
    if FIG.match(raw):
        figure(raw); i += 1; continue
    if TAB_CAP.match(raw):
        flush_alt(); para(s, style='Caption Bang'); i += 1; continue
    if UC_TITLE.match(s):
        flush_alt(); para(s, style='UC Title'); i += 1; continue
    if s.startswith('['):
        mp = PUML.search(s)
        if mp:
            pending['png'], pending['marker'] = mp.group(1), s
            s = s.rstrip(']') + ' – đã chèn ảnh phác thảo render từ file này ngay bên dưới]'
        para(s, style='Ghi chu'); i += 1; continue

    flush_alt()
    if s.startswith('* ') or s.startswith('•'):
        para(re.sub(r'^(\*|•)\s*', '', s), style='List Bullet').alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        para(s, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    i += 1

# Style danh sách gạch đầu dòng dùng cùng font
set_font(st['List Bullet'], 13)
doc.save(OUT)
print('ok', stats)
